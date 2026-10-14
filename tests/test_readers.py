"""Tes pembaca file. Tidak butuh database, hanya openpyxl."""

import datetime as dt
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from apps.imports.readers import (
    MAX_COLS,
    ReaderError,
    build_columns,
    cell_to_text,
    iter_records,
    list_sheets,
    load_table,
)


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.addCleanup(self._tmp.cleanup)

    def xlsx(self, rows, name="data.xlsx", sheet="Data", merge=()):
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row in rows:
            ws.append(row)
        for ref in merge:
            ws.merge_cells(ref)
        path = self.dir / name
        wb.save(path)
        return path


class CellToTextTests(unittest.TestCase):
    def test_nilai_dasar(self):
        self.assertIsNone(cell_to_text(None))
        self.assertIsNone(cell_to_text("   "))
        self.assertEqual(cell_to_text(30.0), "30")
        self.assertEqual(cell_to_text(2.5), "2.5")
        self.assertEqual(cell_to_text(7), "7")
        self.assertEqual(cell_to_text(True), "true")

    def test_tanggal(self):
        self.assertEqual(cell_to_text(dt.datetime(2026, 3, 1)), "2026-03-01")
        self.assertEqual(cell_to_text(dt.datetime(2026, 3, 1, 8, 30)), "2026-03-01T08:30:00")
        self.assertEqual(cell_to_text(dt.date(2026, 3, 1)), "2026-03-01")

    def test_spasi_dirapikan(self):
        self.assertEqual(cell_to_text("  Pra-   Berkembang "), "Pra- Berkembang")


class XlsxTests(Base):
    def test_npsn_teks_dengan_nol_di_depan_tetap_utuh(self):
        path = self.xlsx([["NPSN", "Nama"], ["01234567", "SMP A"], [20100002, "SMP B"]])
        table = load_table(path)
        self.assertEqual(table.rows[1][0], "01234567")
        self.assertEqual(table.rows[2][0], "20100002")

    def test_baris_kosong_di_ujung_dibuang_dan_lebar_diseragamkan(self):
        path = self.xlsx([["A", "B", "C"], ["1", None, None], [None, None, None], []])
        table = load_table(path)
        self.assertEqual(len(table.rows), 2)
        self.assertEqual({len(r) for r in table.rows}, {3})
        self.assertEqual(table.n_cols, 3)

    def test_pilih_sheet(self):
        wb = Workbook()
        wb.active.title = "Satu"
        wb.active.append(["x"])
        wb.create_sheet("Dua").append(["y"])
        path = self.dir / "dua.xlsx"
        wb.save(path)
        self.assertEqual(list_sheets(path), ["Satu", "Dua"])
        self.assertEqual(load_table(path, sheet="Dua").rows, [["y"]])
        with self.assertRaisesRegex(ReaderError, "Pilih nama sheet"):
            load_table(path)
        with self.assertRaisesRegex(ReaderError, "Sheet 'Tiga' tidak ada"):
            load_table(path, sheet="Tiga")

    def test_formula_dibaca_sebagai_nilai_tersimpan(self):
        # data_only=True: file yang dibuat openpyxl tidak punya nilai cache, jadi None
        wb = Workbook()
        wb.active.append([1, 2, "=A1+B1"])
        path = self.dir / "f.xlsx"
        wb.save(path)
        self.assertEqual(load_table(path).rows[0][:2], ["1", "2"])

    def test_tolak_xls_xlsm_dan_ekstensi_lain(self):
        for name, pesan in [("a.xls", ".xls"), ("a.xlsm", "makro"), ("a.pdf", "Hanya file")]:
            path = self.dir / name
            path.write_bytes(b"x")
            with self.assertRaisesRegex(ReaderError, pesan):
                load_table(path)

    def test_tolak_xlsx_palsu(self):
        path = self.dir / "palsu.xlsx"
        path.write_bytes(b"ini bukan zip")
        with self.assertRaisesRegex(ReaderError, "bukan file Excel"):
            load_table(path)

    def test_tolak_zip_tanpa_workbook(self):
        path = self.dir / "zip.xlsx"
        with zipfile.ZipFile(path, "w") as z:
            z.writestr("hello.txt", "hi")
        with self.assertRaisesRegex(ReaderError, "bukan file Excel"):
            load_table(path)

    def test_xlsx_rusak_diberi_pesan_bersahabat(self):
        path = self.dir / "rusak.xlsx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("xl/workbook.xml", "<workbook")
        with self.assertRaisesRegex(ReaderError, "rusak atau tidak dapat dibaca"):
            load_table(path)

    def test_tolak_xlsx_berisi_makro(self):
        path = self.xlsx([["a"]], name="makro.xlsx")
        with zipfile.ZipFile(path, "a") as z:
            z.writestr("xl/vbaProject.bin", b"\x00")
        with self.assertRaisesRegex(ReaderError, "makro"):
            load_table(path)

    def test_tolak_terlalu_banyak_kolom(self):
        path = self.xlsx([[f"k{i}" for i in range(MAX_COLS + 1)]])
        with self.assertRaisesRegex(ReaderError, "kolom"):
            load_table(path)


class CsvTests(Base):
    def csv(self, text, encoding="utf-8", name="d.csv"):
        path = self.dir / name
        path.write_bytes(text.encode(encoding))
        return path

    def test_titik_koma_dan_bom(self):
        path = self.csv("﻿NPSN;Nama\n01234567;SMP A\n")
        table = load_table(path, encoding="utf-8", delimiter=";")
        self.assertEqual(table.rows, [["NPSN", "Nama"], ["01234567", "SMP A"]])

    def test_pemisah_dan_encoding_harus_valid(self):
        path = self.csv("a;b\n")
        with self.assertRaisesRegex(ReaderError, "Pilih encoding dan pemisah"):
            load_table(path)
        with self.assertRaisesRegex(ReaderError, "Pemisah"):
            load_table(path, encoding="utf-8", delimiter="|")
        with self.assertRaisesRegex(ReaderError, "Encoding"):
            load_table(path, encoding="latin-9", delimiter=";")

    def test_salah_encoding_diberi_pesan_ramah(self):
        path = self.csv("Nama\nSMP Café\n", encoding="windows-1252")
        with self.assertRaisesRegex(ReaderError, "encoding lain"):
            load_table(path, encoding="utf-8", delimiter=",")
        table = load_table(path, encoding="windows-1252", delimiter=",")
        self.assertEqual(table.rows[1][0], "SMP Café")


class ColumnTests(Base):
    def test_header_tunggal(self):
        table = load_table(self.xlsx([["NPSN", "Nama"], ["1", "A"]]))
        cols = build_columns(table, [1])
        self.assertEqual([(c.letter, c.header) for c in cols], [("A", "NPSN"), ("B", "Nama")])

    def test_header_bertingkat_dengan_sel_gabungan(self):
        rows = [
            ["Profil", None, "PM", None, None],
            ["NPSN", "Nama", "1.1.1", "1.1.2", "Skor"],
            ["10100001", "SMP A", "3", "4", "3.5"],
        ]
        path = self.xlsx(rows, merge=["A1:B1", "C1:E1"])
        table = load_table(path)
        cols = build_columns(table, [1, 2], fill_right=[1])
        self.assertEqual(
            [c.header for c in cols],
            [
                "Profil | NPSN",
                "Profil | Nama",
                "PM | 1.1.1",
                "PM | 1.1.2",
                "PM | Skor",
            ],
        )

    def test_tanpa_fill_right_sel_gabungan_kosong_hanya_memakai_baris_lain(self):
        rows = [["PM", None], ["1.1.1", "1.1.2"]]
        cols = build_columns(load_table(self.xlsx(rows)), [1, 2])
        self.assertEqual([c.header for c in cols], ["PM | 1.1.1", "1.1.2"])

    def test_header_kembar_dan_kosong(self):
        table = load_table(self.xlsx([["Skor", "Skor", None, "Skor"], ["1", "2", "3", "4"]]))
        cols = build_columns(table, [1])
        self.assertEqual([c.header for c in cols], ["Skor", "Skor (2)", "Kolom C", "Skor (3)"])

    def test_header_di_luar_isi_file(self):
        table = load_table(self.xlsx([["a"]]))
        with self.assertRaisesRegex(ReaderError, "di luar isi file"):
            build_columns(table, [5])

    def test_iter_records_melewati_baris_kosong_dan_memberi_nomor_baris_asli(self):
        rows = [["NPSN", "Nama"], ["1", "A"], [None, None], ["3", "C"]]
        table = load_table(self.xlsx(rows))
        cols = build_columns(table, [1])
        got = list(iter_records(table, cols, 2))
        self.assertEqual([n for n, _ in got], [2, 4])
        self.assertEqual(got[1][1], {"NPSN": "3", "Nama": "C"})


if __name__ == "__main__":
    unittest.main()
