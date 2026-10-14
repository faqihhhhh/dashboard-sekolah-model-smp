"""Ekstrak kamus dan rubrik dari workbook baseline ke CSV aplikasi.

Script hanya menyalin isi sheet codebook dan MASTER, bukan baris data sekolah.
"""

import csv
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests/fixtures/profil_kondisi_awal_baseline_sm_smp.xlsx"
OUTPUT = ROOT / "apps/baseline/data"
RATINGS = {
    "pra berkembang": "pra_berkembang",
    "berkembang": "berkembang",
    "maju": "maju",
    "unggul": "unggul",
}


def _clean(value):
    return " ".join(str(value).replace("\u00a0", " ").split()) if value is not None else ""


def _write_csv(path, headers, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def build_reference(source=SOURCE, output=OUTPUT):
    workbook = load_workbook(source, read_only=False, data_only=True)
    indicators = []
    codebook_by_key = {}
    orders = {
        ("pm", "L1"): 0,
        ("pm", "L2"): 0,
        ("pm", "L3"): 0,
        ("kka", "L1"): 0,
        ("kka", "L2"): 0,
        ("kka", "L3"): 0,
    }

    for framework, sheet_name in (("pm", "codebook PM"), ("kka", "codebook KKA")):
        current_l1 = current_l2 = None
        sheet = workbook[sheet_name]
        for row_number in range(2, sheet.max_row + 1):
            level1, level2, code2, level3, code3 = [
                _clean(sheet.cell(row_number, col).value) for col in range(1, 6)
            ]
            if level1:
                current_l1 = level1
                orders[(framework, "L1")] += 1
                indicators.append(
                    {
                        "kerangka": framework,
                        "level": "L1",
                        "kode": current_l1,
                        "nama": current_l1,
                        "parent_level": "",
                        "parent_code": "",
                        "urutan": orders[(framework, "L1")],
                    }
                )
                codebook_by_key[(framework, "L1", current_l1)] = current_l1
            if code2:
                if not current_l1:
                    raise ValueError(f"{sheet_name} baris {row_number}: Level 2 tanpa Level 1")
                current_l2 = code2
                orders[(framework, "L2")] += 1
                indicators.append(
                    {
                        "kerangka": framework,
                        "level": "L2",
                        "kode": code2,
                        "nama": level2,
                        "parent_level": "L1",
                        "parent_code": current_l1,
                        "urutan": orders[(framework, "L2")],
                    }
                )
                codebook_by_key[(framework, "L2", code2)] = level2
            if code3:
                if not current_l2:
                    raise ValueError(f"{sheet_name} baris {row_number}: Level 3 tanpa Level 2")
                orders[(framework, "L3")] += 1
                indicators.append(
                    {
                        "kerangka": framework,
                        "level": "L3",
                        "kode": code3,
                        "nama": level3,
                        "parent_level": "L2",
                        "parent_code": current_l2,
                        "urutan": orders[(framework, "L3")],
                    }
                )
                codebook_by_key[(framework, "L3", code3)] = level3

    source_to_indicator = {}
    data_sheet = workbook["SMP"]
    framework = None
    for col in range(1, data_sheet.max_column + 1):
        group = _clean(data_sheet.cell(1, col).value)
        if group:
            if group.startswith("PM Level 3"):
                framework = "pm"
            elif group.startswith("KKA Level 3"):
                framework = "kka"
            else:
                framework = None
        code = _clean(data_sheet.cell(2, col).value)
        if framework and code:
            source_to_indicator[get_column_letter(col)] = (framework, code)

    guide_rows = []
    master = workbook["MASTER"]
    for row_number in range(2, master.max_row + 1):
        source_column = _clean(master.cell(row_number, 1).value)
        rating_source = _clean(master.cell(row_number, 2).value)
        description = _clean(master.cell(row_number, 3).value)
        if not source_column and not rating_source and not description:
            continue
        key = " ".join(re.sub(r"[^a-z]+", " ", rating_source.lower()).split())
        if key not in RATINGS:
            raise ValueError(f"MASTER baris {row_number}: rating tidak dikenal")
        if source_column not in source_to_indicator:
            raise ValueError(f"MASTER baris {row_number}: kolom sumber tidak ada di sheet SMP")
        framework, code = source_to_indicator[source_column]
        if (framework, "L3", code) not in codebook_by_key:
            raise ValueError(f"MASTER baris {row_number}: kode indikator tidak ada di codebook")
        guide_rows.append(
            {
                "kerangka": framework,
                "kode_indikator": code,
                "rating": RATINGS[key],
                "deskripsi": description,
            }
        )

    workbook.close()
    output.mkdir(parents=True, exist_ok=True)
    _write_csv(
        output / "indikator.csv",
        ["kerangka", "level", "kode", "nama", "parent_level", "parent_code", "urutan"],
        indicators,
    )
    _write_csv(
        output / "panduan_rating.csv",
        ["kerangka", "kode_indikator", "rating", "deskripsi"],
        guide_rows,
    )
    return len(indicators), len(guide_rows)


if __name__ == "__main__":
    indicator_count, guide_count = build_reference()
    print(f"Selesai: {indicator_count} indikator, {guide_count} panduan rating.")
