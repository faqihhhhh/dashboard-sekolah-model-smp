"""Ekstrak agregat kendala dan dukungan dari workbook baseline."""

import csv
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "tests/fixtures/profil_kondisi_awal_baseline_sm_smp.xlsx"
OUTPUT = ROOT / "apps/intervensi/data/agregat_baseline.csv"


def clean(value):
    return " ".join(str(value).replace("\u00a0", " ").split()) if value is not None else ""


def rating_key(value):
    normalized = re.sub(r"\s+", " ", value.lower()).strip()
    normalized = re.sub(r"\s*[-–]\s*", "-", normalized)
    if normalized.startswith("pra"):
        return "pra_berkembang", "Pra-Berkembang"
    key = normalized.replace(" ", "_")
    labels = {"berkembang": "Berkembang", "maju": "Maju", "unggul": "Unggul"}
    if key not in labels:
        raise ValueError(f"Tingkat baseline tidak dikenal: {value}")
    return key, labels[key]


def parse_blocks(sheet, kind, column):
    rows = []
    for row_number in range(1, sheet.max_row + 1):
        role_heading = clean(sheet.cell(row_number, column).value)
        match = re.fullmatch(r"(Kepala Sekolah|Guru)\s*-\s*(PM|KKA)", role_heading, re.I)
        if not match:
            continue
        target = "kepala_sekolah" if match.group(1).lower().startswith("kepala") else "guru"
        framework = match.group(2).lower()
        header_row = row_number + 1
        header = clean(sheet.cell(header_row, column).value)
        if not header:
            raise ValueError(f"{sheet.title} baris {header_row}: judul kolom kosong")
        groups = []
        for header_column in range(column + 1, sheet.max_column + 1):
            label = clean(sheet.cell(header_row, header_column).value)
            if not label:
                break
            n_match = re.search(r"\(\s*n\s*=\s*(\d+)\s*\)", label, re.I)
            if not n_match:
                message = f"{sheet.title} baris {header_row}: jumlah responden tidak ada: {label}"
                raise ValueError(message)
            level_text = clean(label[: n_match.start()])
            level, level_label = rating_key(level_text)
            groups.append((header_column, level, level_label, int(n_match.group(1))))

        order = 0
        for data_row in range(header_row + 1, sheet.max_row + 1):
            label = clean(sheet.cell(data_row, column).value)
            if not label:
                break
            values = []
            for value_column, level, level_label, sample_size in groups:
                value = sheet.cell(data_row, value_column).value
                if isinstance(value, (float, int)) and not isinstance(value, bool):
                    values.append((level, level_label, sample_size, value))
            if not values:
                break
            order += 1
            for level, level_label, sample_size, proportion in values:
                rows.append(
                    {
                        "jenis": kind,
                        "kerangka": framework,
                        "sasaran": target,
                        "kategori": label,
                        "tingkat": level,
                        "tingkat_label": level_label,
                        "jumlah_responden": sample_size,
                        "proporsi": str(proportion),
                        "urutan": order,
                    }
                )
    return rows


def build_reference(source=SOURCE, output=OUTPUT):
    workbook = load_workbook(source, read_only=False, data_only=True)
    rows = []
    for kind, sheet_name in (("kendala", "Kendala"), ("dukungan", "Dukungan")):
        sheet = workbook[sheet_name]
        rows.extend(parse_blocks(sheet, kind, 1))
        rows.extend(parse_blocks(sheet, kind, 7))
    workbook.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "jenis",
        "kerangka",
        "sasaran",
        "kategori",
        "tingkat",
        "tingkat_label",
        "jumlah_responden",
        "proporsi",
        "urutan",
    ]
    with output.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


if __name__ == "__main__":
    print(f"Selesai: {build_reference()} nilai agregat disimpan ke {OUTPUT.relative_to(ROOT)}")
